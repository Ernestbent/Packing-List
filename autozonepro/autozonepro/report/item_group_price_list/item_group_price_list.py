import os
from io import BytesIO
from urllib.parse import unquote, urlparse

import frappe
from frappe import _
from frappe.desk.utils import provide_binary_file
from frappe.utils import flt


def execute(filters=None):
	models = get_models(filters)
	columns = get_columns(models)
	data = get_data(filters)
	return columns, data


@frappe.whitelist()
def get_model_options(txt=""):
	return frappe.get_all(
		"Item",
		filters=[
			["Item", "disabled", "=", 0],
			["Item", "custom_model", "!=", ""],
			["Item", "custom_model", "like", f"%{txt}%"],
		],
		pluck="custom_model",
		distinct=True,
		order_by="custom_model",
		limit_page_length=50,
	)


def get_models(filters):
	filters = frappe._dict(filters or {})
	conditions = "disabled = 0 and custom_model is not null and custom_model != ''"
	if filters.get("item_group"):
		conditions += " and item_group = %(item_group)s"
	models = get_selected_models(filters)
	if models:
		filters.models = tuple(models)
		conditions += " and custom_model in %(models)s"

	rows = frappe.db.sql("""
		select distinct custom_model
		from `tabItem`
		where {conditions}
		order by custom_model
	""".format(conditions=conditions), filters, as_dict=1)

	return [row.custom_model for row in rows]


def get_selected_models(filters):
	models = filters.get("model")
	if not models:
		return []

	if isinstance(models, str):
		models = models.strip()
		if not models:
			return []
		if models.startswith("["):
			models = frappe.parse_json(models)
		else:
			models = [models]

	return [model for model in models if model]


def get_columns(models):
	columns = [
		{"label": "Group", "fieldname": "group", "fieldtype": "Data", "width": 150},
		{"label": "Sub Group", "fieldname": "sub_group", "fieldtype": "Data", "width": 150},
		{"label": "Brand", "fieldname": "brand", "fieldtype": "Data", "width": 100},
		{"label": _("Image"), "fieldname": "picture", "fieldtype": "HTML", "width": 190},
	]
	for model in models:
		fieldname = frappe.scrub(model)
		columns.append({"label": model, "fieldname": fieldname, "fieldtype": "Currency", "width": 140})
	return columns


def get_data(filters):
	filters = frappe._dict(filters or {})

	## walk the Item Group tree once, same as before, to resolve each item's top-level group
	group_rows = frappe.db.sql("select name, parent_item_group from `tabItem Group`", as_dict=1)
	parent_map = {}
	for row in group_rows:
		parent_map[row.name] = row.parent_item_group

	resolved = {}

	def resolve_group(item_group):
		if item_group in resolved:
			return resolved[item_group]
		path = [item_group]
		current = item_group
		while parent_map.get(current) and parent_map.get(current) != "All Item Groups":
			current = parent_map[current]
			path.append(current)
		top = path[-1]
		resolved[item_group] = top
		return top

	# Items without a model cannot populate a dynamic price column.
	conditions = "disabled = 0 and custom_model is not null and custom_model != ''"
	if filters.get("item_group"):
		conditions += " and item_group = %(item_group)s"
	models = get_selected_models(filters)
	if models:
		filters.models = tuple(models)
		conditions += " and custom_model in %(models)s"

	items = frappe.db.sql("""
		select item_group, brand, custom_model, standard_rate, image
		from `tabItem`
		where {conditions}
		order by item_group, brand, custom_model, item_code
	""".format(conditions=conditions), filters, as_dict=1)

	# One row per brand; each corresponding model fills a price column on that row.
	rows = {}
	for item in items:
		top = resolve_group(item.item_group)
		if top == "Products":
			continue

		sub = "" if item.item_group == top else item.item_group
		key = (top, sub, item.brand)

		if key not in rows:
			rows[key] = {
				"group": top,
				"sub_group": sub,
				"brand": item.brand,
				"picture": None,
			}

		fieldname = frappe.scrub(item.custom_model)
		rows[key][fieldname] = item.standard_rate
		if item.image and not rows[key]["picture"]:
			rows[key]["picture"] = item.image

	# Sort by the resolved first column before the child fields. This keeps every
	# subgroup/brand belonging to a Group together until that Group is exhausted.
	return sorted(
		rows.values(),
		key=lambda row: (
			(row.get("group") or "").casefold(),
			(row.get("sub_group") or "").casefold(),
			(row.get("brand") or "").casefold(),
		),
	)


@frappe.whitelist()
def download_xlsx_with_images(filters=None):
	filters = frappe._dict(frappe.parse_json(filters) if filters else {})
	columns, data = execute(filters)
	content = build_xlsx_with_images(columns, data, _("Item Group Price List"))
	provide_binary_file("Item Group Price List With Images", "xlsx", content)


def build_xlsx_with_images(columns, data, title):
	from openpyxl import Workbook
	from openpyxl.drawing.image import Image as ExcelImage
	from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
	from openpyxl.utils import get_column_letter

	wb = Workbook()
	ws = wb.active
	ws.title = "Price List"

	header_fill = PatternFill("solid", fgColor="EDEDED")
	title_fill = PatternFill("solid", fgColor="F7FAFC")
	border = Border(
		left=Side(style="thin", color="D9D9D9"),
		right=Side(style="thin", color="D9D9D9"),
		top=Side(style="thin", color="D9D9D9"),
		bottom=Side(style="thin", color="D9D9D9"),
	)
	center = Alignment(horizontal="center", vertical="center", wrap_text=True)
	left = Alignment(horizontal="left", vertical="center", wrap_text=True)
	right = Alignment(horizontal="right", vertical="center", wrap_text=True)

	ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(columns))
	title_cell = ws.cell(row=1, column=1, value=title)
	title_cell.font = Font(bold=True, size=14)
	title_cell.alignment = center
	title_cell.fill = title_fill
	title_cell.border = border
	ws.row_dimensions[1].height = 26

	for col_index, column in enumerate(columns, start=1):
		cell = ws.cell(row=2, column=col_index, value=column["label"])
		cell.font = Font(bold=True)
		cell.fill = header_fill
		cell.alignment = center
		cell.border = border
		ws.column_dimensions[get_column_letter(col_index)].width = get_excel_column_width(
			column.get("width")
		)

	image_streams = []
	picture_column = get_column_index(columns, "picture")
	if picture_column:
		ws.column_dimensions[get_column_letter(picture_column)].width = pixels_to_excel_width(96)

	for row_index, row in enumerate(data, start=3):
		ws.row_dimensions[row_index].height = pixels_to_points(80)

		for col_index, column in enumerate(columns, start=1):
			fieldname = column["fieldname"]
			value = "" if fieldname == "picture" else row.get(fieldname)
			cell = ws.cell(row=row_index, column=col_index, value=value)
			cell.border = border
			cell.alignment = left

			if column.get("fieldtype") == "Currency":
				cell.alignment = right
				cell.number_format = '#,##0'
			elif fieldname == "picture":
				cell.alignment = center

		image_url = row.get("picture")
		if image_url and picture_column:
			try:
				image_source = get_excel_image_source(image_url)
				if image_source:
					if not isinstance(image_source, str):
						image_streams.append(image_source)

					excel_image = ExcelImage(image_source)
					# Scale the Excel drawing while retaining the original embedded image bytes.
					scale = min(80 / excel_image.width, 60 / excel_image.height, 1)
					excel_image.width *= scale
					excel_image.height *= scale
					ws.add_image(excel_image, f"{get_column_letter(picture_column)}{row_index}")
			except Exception:
				log_image_export_error()

	# Keep image buffers alive until the workbook has been saved.
	wb._autozonepro_image_streams = image_streams

	xlsx_file = BytesIO()
	wb.save(xlsx_file)
	return xlsx_file.getvalue()


def get_excel_column_width(report_width):
	if not report_width:
		return 15
	return max(10, min(45, flt(report_width) / 7))


def pixels_to_points(pixels):
	return flt(pixels) * 0.75


def pixels_to_excel_width(pixels):
	return flt(pixels) / 7


def get_column_index(columns, fieldname):
	for index, column in enumerate(columns, start=1):
		if column.get("fieldname") == fieldname:
			return index
	return None


def get_excel_image_source(image_url):
	image_path = get_local_image_path(image_url)

	if image_path and os.path.exists(image_path):
		return image_path

	remote_image = get_remote_image_content(image_url)
	if remote_image:
		return remote_image

	return None


def log_image_export_error():
	try:
		frappe.log_error(
			title="Item Group Price List Image Export Failed",
			message=frappe.get_traceback(),
		)
	except Exception:
		pass


def get_local_image_path(image_url):
	if not image_url:
		return None

	parsed_url = urlparse(image_url)
	if parsed_url.scheme and parsed_url.netloc:
		site_url = urlparse(frappe.utils.get_url())
		if parsed_url.netloc != site_url.netloc:
			return None
		image_url = parsed_url.path

	image_url = unquote(image_url.split("?", 1)[0])

	if image_url.startswith("/private/files/"):
		return frappe.get_site_path(image_url.lstrip("/"))
	if image_url.startswith("/files/"):
		return frappe.get_site_path("public", image_url.lstrip("/"))

	file_name = frappe.db.get_value("File", {"file_url": image_url}, "name")
	if file_name:
		return frappe.get_doc("File", file_name).get_full_path()

	return None


def get_remote_image_content(image_url):
	parsed_url = urlparse(image_url or "")
	if parsed_url.scheme not in ("http", "https"):
		return None

	import requests

	try:
		response = requests.get(image_url, timeout=10)
		response.raise_for_status()
	except Exception:
		return None

	if len(response.content) > 5 * 1024 * 1024:
		return None

	return BytesIO(response.content)
